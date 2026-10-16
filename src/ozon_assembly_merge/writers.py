"""Сохранение результата: XLSX, CSV, PDF."""

from __future__ import annotations

import csv
from importlib import resources
from pathlib import Path

from .merge import MergedItem, MergeResult

HEADERS = ["№", "Товар", "Артикул", "Кол-во (итого)", "Источники (склад: кол-во)"]
FORMATS = ("xlsx", "csv", "pdf")


def _rows(items: list[MergedItem]):
    for n, item in enumerate(items, start=1):
        yield [n, item.product, item.article, item.qty, item.sources_text()]


def write_xlsx(result: MergeResult, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    merged_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [6, 70, 38, 14, 40]

    def fill_sheet(ws, items: list[MergedItem]) -> int:
        ws.append(HEADERS)
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = header_fill
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r, (item, values) in enumerate(zip(items, _rows(items)), start=2):
            ws.append(values)
            for c in range(1, len(HEADERS) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=c in (2, 5))
                if item.is_merged:
                    cell.fill = merged_fill
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return len(items) + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Объединённый список"
    last = fill_sheet(ws, result.items)
    ws.auto_filter.ref = f"A1:E{last}"
    ws.cell(row=last + 1, column=2, value="ИТОГО:").font = Font(bold=True)
    ws.cell(row=last + 1, column=4, value=result.total_qty).font = Font(bold=True)
    fill_sheet(wb.create_sheet("Только объединённые"), result.merged_items)
    wb.save(path)


def write_csv(result: MergeResult, path: Path, delimiter: str = ";") -> None:
    # utf-8-sig и «;» — чтобы Excel с русской локалью открыл файл без мастера импорта
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow(HEADERS)
        w.writerows(_rows(result.items))


def _font_path() -> str:
    return str(resources.files("ozon_assembly_merge") / "fonts" / "DejaVuSans.ttf")


def write_pdf(result: MergeResult, path: Path, title: str = "Сборочный лист (объединённый)") -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = "DejaVuSans"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, _font_path()))

    cell = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10)
    head = ParagraphStyle("head", fontName=font, fontSize=8, leading=10, textColor=colors.white)
    h1 = ParagraphStyle("h1", fontName=font, fontSize=13, leading=16)
    meta = ParagraphStyle("meta", fontName=font, fontSize=9, leading=12)

    data = [[Paragraph(h, head) for h in ["№", "Товар", "Артикул", "Кол-во", "Источники"]]]
    for values in _rows(result.items):
        data.append([Paragraph(_escape(str(v)), cell) for v in values])
    data.append(["", Paragraph("ИТОГО", cell), "", Paragraph(str(result.total_qty), cell), ""])

    width = A4[0] - 24 * mm
    table = Table(data, colWidths=[0.05 * width, 0.38 * width, 0.25 * width, 0.08 * width, 0.24 * width], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for r, item in enumerate(result.items, start=1):
        if item.is_merged:
            style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#FFF2CC")))
    table.setStyle(TableStyle(style))

    info = (
        f"Дата: {', '.join(result.dates) or '—'} · позиций: {len(result.items)} · "
        f"объединено: {len(result.merged_items)} · всего штук: {result.total_qty}"
    )
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    doc.build([Paragraph(title, h1), Paragraph(info, meta), Spacer(1, 4 * mm), table])


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write(result: MergeResult, path: str | Path, fmt: str | None = None) -> Path:
    path = Path(path)
    fmt = (fmt or path.suffix.lstrip(".")).lower()
    if fmt not in FORMATS:
        raise ValueError(f"неизвестный формат «{fmt}», поддерживаются: {', '.join(FORMATS)}")
    {"xlsx": write_xlsx, "csv": write_csv, "pdf": write_pdf}[fmt](result, path)
    return path
