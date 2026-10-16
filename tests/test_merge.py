from __future__ import annotations

import csv

import pdfplumber
import pytest
from openpyxl import load_workbook

from ozon_assembly_merge import AssemblyParseError, merge_lists, parse_assembly_list, write
from ozon_assembly_merge.cli import main

from .synthetic import CONT_PAGE_TOP, FIRST_PAGE_TABLE_TOP, Item, catalog, make_assembly_pdf


def _name(seller: int, n: int = 1) -> str:
    return f"carriage_assembly_list_seller_{seller}_time_2026-10-17_{n}.pdf"


@pytest.fixture
def two_lists(tmp_path):
    a = make_assembly_pdf(tmp_path / _name(100001), catalog(30), warehouse="Склад Север")
    b = make_assembly_pdf(tmp_path / _name(100002), catalog(10, offset=2, qty_base=3), warehouse="Склад Юг")
    return a, b


def test_parses_every_row_across_pages(tmp_path):
    items = catalog(30)
    pdf = make_assembly_pdf(tmp_path / _name(100001), items, shipments=21)
    with pdfplumber.open(pdf) as doc:
        assert len(doc.pages) == 3

    parsed = parse_assembly_list(pdf)

    assert [r.num for r in parsed.rows] == list(range(1, 31))
    assert [(r.product, r.article, r.qty) for r in parsed.rows] == [(i.product, i.article, i.qty) for i in items]
    assert parsed.date == "17.10.2026"
    assert parsed.warehouse == "Тестовый склад"
    assert parsed.seller == "100001"
    assert parsed.shipments == 21


def test_first_row_of_continuation_page_sits_above_first_line(tmp_path):
    """Регрессия: первая строка страницы-продолжения начинается выше первой
    линии страницы; без виртуальной границы 0.0 она терялась молча."""
    pdf = make_assembly_pdf(tmp_path / _name(100001), catalog(14))
    with pdfplumber.open(pdf) as doc:
        page2 = doc.pages[1]
        first_line = min(ln["top"] for ln in page2.lines)
        anchor = next(w for w in page2.extract_words() if w["text"] == "13")
        assert anchor["top"] < first_line
        assert first_line > CONT_PAGE_TOP
    assert FIRST_PAGE_TABLE_TOP > CONT_PAGE_TOP

    rows = parse_assembly_list(pdf).rows
    assert [r.num for r in rows] == list(range(1, 15))


def test_gap_in_numbering_is_an_error(tmp_path):
    pdf = make_assembly_pdf(tmp_path / _name(100001), catalog(3), numbers=[1, 2, 4])
    with pytest.raises(AssemblyParseError, match=r"пропущены \[3\]"):
        parse_assembly_list(pdf)


def test_merge_sums_by_article_not_by_name(two_lists):
    lists = [parse_assembly_list(p) for p in two_lists]
    result = merge_lists(lists)

    assert result.input_rows == 40
    assert result.total_qty == sum(a.total_qty for a in lists)
    by_art = {i.article: i for i in result.items}
    assert len(by_art) == len(result.items)

    blue, pink = by_art["WATCH-KIDS-BLUE"], by_art["WATCH-KIDS-PINK"]
    assert blue.product == pink.product
    assert blue.is_merged and pink.is_merged

    expected = sum(r.qty for a in lists for r in a.rows if r.article == "WATCH-KIDS-BLUE")
    assert blue.qty == expected
    assert [label for label, _ in blue.sources] == ["Склад Север (seller 100001)", "Склад Юг (seller 100002)"]
    assert len(result.merged_items) == 10
    assert result.warnings == []


def test_warns_on_same_article_with_different_name(tmp_path):
    a = make_assembly_pdf(tmp_path / _name(1), [Item("Кружка белая", "MUG-1", 1)])
    b = make_assembly_pdf(tmp_path / _name(2), [Item("Кружка белая 300 мл", "MUG-1", 2)], date="18.10.2026")
    result = merge_lists([parse_assembly_list(a), parse_assembly_list(b)])
    assert result.items[0].qty == 3
    assert any("разными названиями" in w for w in result.warnings)
    assert any("по-разному" in w for w in result.warnings)


def test_sort_by_qty(two_lists):
    result = merge_lists([parse_assembly_list(p) for p in two_lists], sort="qty")
    qtys = [i.qty for i in result.items]
    assert qtys == sorted(qtys, reverse=True)


def test_writers(two_lists, tmp_path):
    result = merge_lists([parse_assembly_list(p) for p in two_lists])

    xlsx = write(result, tmp_path / "out.xlsx")
    wb = load_workbook(xlsx)
    ws = wb["Объединённый список"]
    assert ws.cell(row=1, column=3).value == "Артикул"
    assert ws.cell(row=ws.max_row, column=2).value == "ИТОГО:"
    assert ws.cell(row=ws.max_row, column=4).value == result.total_qty
    assert wb["Только объединённые"].max_row == len(result.merged_items) + 1

    out_csv = write(result, tmp_path / "out.csv")
    with open(out_csv, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert len(rows) == len(result.items) + 1
    assert sum(int(r[3]) for r in rows[1:]) == result.total_qty

    out_pdf = write(result, tmp_path / "out.pdf")
    with pdfplumber.open(out_pdf) as doc:
        text = "\n".join(p.extract_text() or "" for p in doc.pages)
    assert "WATCH-KIDS-PINK" in text
    assert f"всего штук: {result.total_qty}" in text

    with pytest.raises(ValueError):
        write(result, tmp_path / "out.txt")


def test_cli_folder_and_multiple_outputs(two_lists, tmp_path, capsys):
    folder = two_lists[0].parent
    out_x, out_p = tmp_path / "r.xlsx", tmp_path / "r.pdf"
    assert main([str(folder), "-o", str(out_x), "-o", str(out_p)]) == 0
    assert out_x.exists() and out_p.exists()
    out = capsys.readouterr().out
    assert "Строк на входе: 40" in out


def test_cli_default_output_name(tmp_path):
    pdf = make_assembly_pdf(tmp_path / _name(1), catalog(3))
    assert main([str(pdf), "-q", "-f", "csv"]) == 0
    assert (tmp_path / "Сборочный_лист_объединённый_17-10-2026.csv").exists()


def test_cli_strict_dates(tmp_path):
    a = make_assembly_pdf(tmp_path / _name(1), catalog(2))
    b = make_assembly_pdf(tmp_path / _name(2), catalog(2), date="18.10.2026")
    assert main([str(a), str(b), "--strict-dates", "-o", str(tmp_path / "x.xlsx")]) == 3
    assert not (tmp_path / "x.xlsx").exists()


def test_cli_parse_error_exit_code(tmp_path, capsys):
    pdf = make_assembly_pdf(tmp_path / _name(1), catalog(3), numbers=[1, 3, 4])
    assert main([str(pdf), "-o", str(tmp_path / "x.xlsx")]) == 2
    assert "не сохранён" in capsys.readouterr().err
